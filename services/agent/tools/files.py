import os
import re
from pathlib import Path
from typing import Dict, Any, List, Optional

BASE_DIR = Path(__file__).resolve().parent.parent.parent.parent
WORKSPACES_DIR = BASE_DIR / "workspaces"

def _get_workspace(job_id: str) -> Path:
    p = WORKSPACES_DIR / job_id
    p.mkdir(parents=True, exist_ok=True)
    return p

def _normalize_spec(spec_or_content: Any, default_title: str = "Artifact") -> Dict[str, Any]:
    """Ensure spec is a dictionary with title, sections, and references."""
    if isinstance(spec_or_content, dict):
        title = spec_or_content.get("title") or default_title
        sections = spec_or_content.get("sections") or []
        references = spec_or_content.get("references") or []
        if not sections and "content" in spec_or_content:
            sections = [{"heading": "Overview", "content": str(spec_or_content["content"])}]
        return {
            "title": title,
            "sections": sections,
            "references": references
        }
    # Raw string fallback
    content = str(spec_or_content or "")
    lines = content.splitlines()
    title = default_title
    if lines and lines[0].startswith("# "):
        title = lines[0].replace("# ", "").strip()
        content = "\n".join(lines[1:]).strip()
    return {
        "title": title,
        "sections": [{"heading": "Content", "content": content}],
        "references": []
    }

def render_markdown(spec: Dict[str, Any], filepath: Path) -> int:
    """Renders structured specification to a clean Markdown file."""
    lines = [f"# {spec.get('title', 'Document')}\n"]
    for s in spec.get("sections", []):
        heading = s.get("heading") or "Section"
        lines.append(f"## {heading}\n")
        lines.append(f"{s.get('content', '').strip()}\n")

    refs = spec.get("references", [])
    if refs:
        lines.append("## References\n")
        for r in refs:
            lines.append(f"- {r}")
        lines.append("")

    content = "\n".join(lines)
    filepath.write_text(content, encoding="utf-8")
    return filepath.stat().st_size

def render_docx(spec: Dict[str, Any], filepath: Path) -> int:
    """Renders structured specification to a Microsoft Word .docx file."""
    import docx
    doc = docx.Document()
    doc.add_heading(spec.get("title", "Document"), level=1)

    for s in spec.get("sections", []):
        doc.add_heading(s.get("heading", "Section"), level=2)
        body = s.get("content", "").strip()
        for p in body.split("\n\n"):
            p_strip = p.strip()
            if p_strip.startswith(("- ", "* ", "• ")):
                for item in p_strip.splitlines():
                    clean_item = re.sub(r"^[-*•]\s*", "", item).strip()
                    doc.add_paragraph(clean_item, style="List Bullet")
            elif p_strip:
                doc.add_paragraph(p_strip)

    refs = spec.get("references", [])
    if refs:
        doc.add_heading("References", level=2)
        for r in refs:
            doc.add_paragraph(str(r), style="List Bullet")

    doc.save(str(filepath))
    return filepath.stat().st_size

def render_pdf(spec: Dict[str, Any], filepath: Path) -> int:
    """Renders structured specification to a styled PDF using PyMuPDF."""
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz

    doc = fitz.open()
    page = doc.new_page(width=595, height=842) # A4 format

    margin_left = 54
    margin_right = 541
    y = 54

    def add_page_if_needed(space_needed=40):
        nonlocal page, y
        if y + space_needed > 780:
            page = doc.new_page(width=595, height=842)
            y = 54

    # Title
    title = spec.get("title", "Document")
    rect_title = fitz.Rect(margin_left, y, margin_right, y + 40)
    page.insert_textbox(rect_title, title, fontsize=18, fontname="helv", color=(0.1, 0.2, 0.4))
    y += 50

    # Sections
    for s in spec.get("sections", []):
        add_page_if_needed(60)
        heading = s.get("heading", "Section")
        rect_h = fitz.Rect(margin_left, y, margin_right, y + 25)
        page.insert_textbox(rect_h, heading, fontsize=13, fontname="helv", color=(0.15, 0.25, 0.45))
        y += 28

        body = s.get("content", "").strip()
        paragraphs = [p.strip() for p in body.split("\n\n") if p.strip()]
        for p in paragraphs:
            # Estimate height roughly by character length
            approx_height = max(25, int(len(p) / 60 * 14) + 15)
            add_page_if_needed(approx_height)
            rect_p = fitz.Rect(margin_left, y, margin_right, y + approx_height)
            page.insert_textbox(rect_p, p, fontsize=10.5, fontname="helv", color=(0.2, 0.2, 0.2))
            y += approx_height + 10

    # References
    refs = spec.get("references", [])
    if refs:
        add_page_if_needed(50)
        rect_ref_h = fitz.Rect(margin_left, y, margin_right, y + 25)
        page.insert_textbox(rect_ref_h, "References", fontsize=13, fontname="helv", color=(0.15, 0.25, 0.45))
        y += 28
        for r in refs:
            add_page_if_needed(25)
            rect_r = fitz.Rect(margin_left + 10, y, margin_right, y + 20)
            page.insert_textbox(rect_r, f"• {r}", fontsize=9.5, fontname="helv", color=(0.3, 0.3, 0.3))
            y += 20

    doc.save(str(filepath))
    doc.close()
    return filepath.stat().st_size


# --- Tool Handlers ---

def create_markdown(job_id: str, filename: str, spec: Any) -> Dict[str, Any]:
    ws = _get_workspace(job_id)
    fname = filename if filename.endswith(".md") else f"{filename}.md"
    fpath = ws / fname
    normalized = _normalize_spec(spec, default_title=fname.replace(".md", ""))
    size = render_markdown(normalized, fpath)
    return {
        "ok": True,
        "name": fname,
        "path": f"workspaces/{job_id}/{fname}",
        "size": size,
        "type": "file",
        "format": "markdown",
        "title": normalized.get("title"),
        "sections_count": len(normalized.get("sections", []))
    }

def create_docx(job_id: str, filename: str, spec: Any) -> Dict[str, Any]:
    ws = _get_workspace(job_id)
    fname = filename if filename.endswith(".docx") else f"{filename}.docx"
    fpath = ws / fname
    normalized = _normalize_spec(spec, default_title=fname.replace(".docx", ""))
    size = render_docx(normalized, fpath)
    return {
        "ok": True,
        "name": fname,
        "path": f"workspaces/{job_id}/{fname}",
        "size": size,
        "type": "file",
        "format": "docx",
        "title": normalized.get("title"),
        "sections_count": len(normalized.get("sections", []))
    }

def create_pdf(job_id: str, filename: str, spec: Any) -> Dict[str, Any]:
    ws = _get_workspace(job_id)
    fname = filename if filename.endswith(".pdf") else f"{filename}.pdf"
    fpath = ws / fname
    normalized = _normalize_spec(spec, default_title=fname.replace(".pdf", ""))
    size = render_pdf(normalized, fpath)
    return {
        "ok": True,
        "name": fname,
        "path": f"workspaces/{job_id}/{fname}",
        "size": size,
        "type": "file",
        "format": "pdf",
        "title": normalized.get("title"),
        "sections_count": len(normalized.get("sections", []))
    }

def read_generated(job_id: str, filename: str) -> Dict[str, Any]:
    ws = _get_workspace(job_id)
    fpath = ws / filename
    if not fpath.exists():
        return {"ok": False, "error": f"File '{filename}' does not exist in workspace."}
    try:
        if filename.endswith(".md") or filename.endswith(".txt"):
            text = fpath.read_text(encoding="utf-8")
            return {"ok": True, "filename": filename, "content": text[:1500], "size": fpath.stat().st_size}
        else:
            return {"ok": True, "filename": filename, "size": fpath.stat().st_size, "exists": True}
    except Exception as e:
        return {"ok": False, "error": str(e)}
